import csv
import io
import zipfile
from datetime import date, timedelta
from unittest import mock

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.kawasan.tests import make_kawasan
from apps.laporan.constants import Status
from apps.laporan.models import Laporan

from . import queries, services
from .models import Hotspot, HotspotSyncLog
from .sipongi import build_params, extract_features, parse_feature


def feat(hs_id, lon, lat, conf="high", waktu="2026-09-10 10:00", sat="NASA-SNPP"):
    return {"type": "Feature", "geometry": {"type": "Point", "coordinates": [lon, lat]},
            "properties": {"hs_id": hs_id, "confidence_level": conf, "date_hotspot": waktu, "satelit": sat}}


class SipongiParsingTests(TestCase):
    def test_extract_shapes(self):
        f = [feat("a", 1, 1)]
        self.assertEqual(extract_features({"features": {"features": f}}), f)
        self.assertEqual(extract_features({"features": f}), f)
        self.assertEqual(extract_features(f), f)
        self.assertEqual(extract_features({"foo": 1}), [])

    def test_parse_feature(self):
        r = parse_feature(feat("2026-09-10-abc", 110.5, -2.5, "MEDIUM"))
        self.assertEqual((r["tanggal"], r["confidence"], r["lon"], r["lat"]), (date(2026, 9, 10), "medium", 110.5, -2.5))
        self.assertIsNone(parse_feature({"geometry": {"coordinates": ["x", 1]}, "properties": {}}))
        self.assertIsNone(parse_feature(feat("", 1, 1)))
        self.assertIsNone(parse_feature(feat("zzz", 1, 1, waktu="")))
        self.assertIsNone(parse_feature({}))

    def test_date_falls_back_to_waktu(self):
        r = parse_feature(feat("noid", 1, 1, waktu="2026-09-11 01:00"))
        self.assertEqual(r["tanggal"], date(2026, 9, 11))

    def test_params_match_legacy_query(self):
        p = dict(build_params(date(2026, 9, 1), date(2026, 9, 2)))
        self.assertEqual((p["wilayah"], p["filterperiode"], p["from"], p["to"], p["late"]), ("IN", "true", "2026-09-01", "2026-09-02", "custom"))
        self.assertEqual([v for k, v in build_params() if k == "satelit[]"], ["NASA-MODIS", "NASA-SNPP", "NASA-NOAA20"])
        self.assertEqual(dict(build_params())["filterperiode"], "false")


class SyncTests(TestCase):
    def setUp(self):
        self.k = make_kawasan("TN A", "BKSDA X", box=(110, -3, 111, -2))

    def test_sync_stores_matches_and_filters_low(self):
        fetcher = lambda a, b: [feat("h1", 110.5, -2.5), feat("h2", 130, 0, "medium"), feat("h3", 110.6, -2.6, "low"), {"bad": 1}]
        log = services.sync_range(date(2026, 9, 10), date(2026, 9, 10), fetcher)
        self.assertTrue(log.success)
        self.assertEqual((log.fetched, log.stored, log.inside_kawasan), (4, 2, 1))
        self.assertEqual(Hotspot.objects.get(hs_id="h1").kawasan, self.k)
        self.assertIsNone(Hotspot.objects.get(hs_id="h2").kawasan)
        self.assertFalse(Hotspot.objects.filter(hs_id="h3").exists())

    def test_sync_is_idempotent_and_updates(self):
        services.sync_range(None, None, lambda a, b: [feat("h1", 110.5, -2.5, "medium")])
        services.sync_range(None, None, lambda a, b: [feat("h1", 110.5, -2.5, "high"), feat("h1", 110.5, -2.5, "high")])
        self.assertEqual(Hotspot.objects.count(), 1)
        self.assertEqual(Hotspot.objects.get().confidence, "high")

    def test_failure_is_logged_and_raised(self):
        def boom(a, b):
            raise RuntimeError("API mati")
        with self.assertRaises(RuntimeError):
            services.sync_range(None, None, boom)
        log = HotspotSyncLog.objects.get()
        self.assertFalse(log.success)
        self.assertIn("API mati", log.message)

    def test_sync_days_makes_one_request_per_day(self):
        calls = []
        services.sync_days(2, lambda a, b: calls.append((a, b)) or [])
        self.assertEqual(len(calls), 3)
        self.assertEqual(calls[-1][0], timezone.localdate())

    def test_rematch_after_boundary_change(self):
        services.sync_range(None, None, lambda a, b: [feat("h1", 110.5, -2.5)])
        self.k.geom = {"type": "MultiPolygon", "coordinates": [[[[130, 0], [131, 0], [131, 1], [130, 1], [130, 0]]]]}
        self.k.save()
        total, changed = services.rematch_all()
        self.assertEqual((total, changed), (1, 1))
        self.assertIsNone(Hotspot.objects.get().kawasan)


class QueryAndViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.a = make_kawasan("TN A", "UPT 1", box=(110, -3, 111, -2))
        cls.b = make_kawasan("TN B", "UPT 2", box=(120, 0, 121, 1))
        cls.today = timezone.localdate()
        mk = lambda i, kaw, conf, d=0: Hotspot.objects.create(hs_id=f"x{i}", tanggal=cls.today - timedelta(days=d), confidence=conf,
                                                             lat=-2.5, lon=110.5, kawasan=kaw, satelit="S")
        mk(1, cls.a, "high"); mk(2, cls.a, "medium"); mk(3, cls.b, "high", 1); mk(4, None, "high"); mk(5, cls.a, "low")
        mk(6, cls.a, "high", 3)
        cls.user = User.objects.create_user("u", password="pw12345678")
        cls.staff = User.objects.create_user("s", password="pw12345678", is_staff=True)

    def test_summary_counts_only_inside_high_medium(self):
        d0, d1 = queries.default_range()
        self.assertEqual(queries.summary(d0, d1), {"high": 2, "medium": 1, "total": 3})

    def test_top_and_weekly(self):
        d0 = self.today - timedelta(days=7)
        top = queries.top_kawasan(d0, self.today)
        self.assertEqual((top[0]["nkws"], top[0]["total"]), ("TN A", 3))
        w = queries.weekly_by_upt(self.today)
        self.assertEqual(len(w["dates"]), 7)
        self.assertEqual(w["series"]["UPT 1"]["TN A"][-1], 2)
        self.assertEqual(w["series"]["UPT 1"]["TN A"][-4], 1)

    def test_unreported_excludes_kawasan_with_report(self):
        d0, d1 = queries.default_range()
        self.assertEqual({u["nkws"] for u in queries.unreported_kawasan(d0, d1)}, {"TN A", "TN B"})
        Laporan.objects.create(tanggal=self.today, kawasan=self.a, status=Status.UPAYA, lat=-2.5, lon=110.5, luas=0)
        self.assertEqual({u["nkws"] for u in queries.unreported_kawasan(d0, d1)}, {"TN B"})

    def test_parse_range_defaults_and_swap(self):
        self.assertEqual(queries.parse_range({}), queries.default_range())
        self.assertEqual(queries.parse_range({"dari": "2026-09-10", "sampai": "2026-09-01"}), (date(2026, 9, 1), date(2026, 9, 10)))
        self.assertEqual(queries.parse_range({"dari": "bukan-tanggal"}), queries.default_range())

    def test_api_hotspot_and_stats(self):
        self.client.force_login(self.user)
        j = self.client.get(reverse("api_hotspot")).json()
        self.assertEqual(j["summary"]["total"], 3)
        self.assertEqual(len(j["results"]), 4, "termasuk hotspot di luar kawasan, tanpa 'low'")
        self.assertEqual(sum(1 for r in j["results"] if r["kid"] is None), 1)
        s = self.client.get(reverse("api_hotspot_stats")).json()
        self.assertEqual(s["top"][0]["nkws"], "TN A")

    def test_exports(self):
        self.client.force_login(self.user)
        q = f"?dari={self.today - timedelta(days=7)}&sampai={self.today}"
        r = self.client.get(reverse("hotspot_export", args=["csv"]) + q)
        rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
        self.assertEqual(rows[0][:2], ["HS_ID", "Confidence"])
        self.assertEqual(len(rows), 1 + 4)  # dalam kawasan saja: x1,x2,x3,x6
        rows_all = list(csv.reader(io.StringIO(self.client.get(reverse("hotspot_export", args=["csv"]) + q + "&scope=semua").content.decode("utf-8-sig"))))
        self.assertEqual(len(rows_all), 1 + 5)
        r = self.client.get(reverse("hotspot_export", args=["xlsx"]) + q)
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws.max_row, 5)
        r = self.client.get(reverse("hotspot_export", args=["shp"]) + q)
        z = zipfile.ZipFile(io.BytesIO(r.content))
        self.assertEqual(sorted(n.rsplit(".", 1)[1] for n in z.namelist()), ["dbf", "prj", "shp", "shx"])
        self.assertEqual(self.client.get(reverse("hotspot_export", args=["pdf"])).status_code, 404)

    def test_sync_endpoint_staff_only(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.post(reverse("hotspot_sync")).status_code, 403)
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(reverse("hotspot_sync")).status_code, 405)
        with mock.patch("apps.hotspot.services.sync_days", side_effect=RuntimeError("x")):
            self.assertEqual(self.client.post(reverse("hotspot_sync")).status_code, 502)
